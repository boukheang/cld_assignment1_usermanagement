import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

SCREENSHOTS_DIR = r"D:\Agent\cld\University_User_Management_Platform\screenshots"
DOCX_OUT_ROOT = r"D:\Agent\cld\Role_Based_User_Management_Submission.docx"
DOCX_OUT_PROJ = r"D:\Agent\cld\University_User_Management_Platform\Role_Based_User_Management_Submission.docx"
PDF_OUT_ROOT = r"D:\Agent\cld\Role_Based_User_Management_Submission.pdf"
PDF_OUT_PROJ = r"D:\Agent\cld\University_User_Management_Platform\Role_Based_User_Management_Submission.pdf"

doc = docx.Document()

# Set page margins to 0.5 inches
for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

def set_run_font(run, font_name="Calibri", size_pt=11, bold=False, color_rgb=(0,0,0), italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_header_block():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    r1 = p.add_run("Fall 2026 Semester\n")
    set_run_font(r1, "Calibri", 12, bold=True, color_rgb=(80, 80, 80))
    
    r2 = p.add_run("CLD 376 Cloud Native Development\n")
    set_run_font(r2, "Calibri", 14, bold=True, color_rgb=(0, 51, 102))

    p_student = doc.add_paragraph()
    p_student.paragraph_format.space_before = Pt(4)
    p_student.paragraph_format.space_after = Pt(10)
    r3 = p_student.add_run("Goat Ronaldo\n")
    set_run_font(r3, "Calibri", 13, bold=True, color_rgb=(0, 0, 0))
    r4 = p_student.add_run("ID: 2024476\n")
    set_run_font(r4, "Calibri", 11, bold=False, color_rgb=(60, 60, 60))

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(14)
    r5 = p_title.add_run("01. Role-Based User Management System Using Microservices API Gateway NodeJS MongoDB")
    set_run_font(r5, "Calibri", 16, bold=True, color_rgb=(0, 70, 140))

def add_callout_box(text, bold_prefix="Scenario: "):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(7.5)
    
    shading = parse_xml(r'<w:shd {} w:fill="F0F4F8"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)
    
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="0055A5"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r_bold = p.add_run(bold_prefix)
    set_run_font(r_bold, "Calibri", 10.5, bold=True, color_rgb=(0, 51, 102))
    r_body = p.add_run(text)
    set_run_font(r_body, "Calibri", 10.5, bold=False, color_rgb=(40, 40, 40))
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, "Calibri", 13.5, bold=True, color_rgb=(0, 51, 102))

def add_body_p(text, bold_lead=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    if bold_lead:
        r_lead = p.add_run(bold_lead + " ")
        set_run_font(r_lead, "Calibri", 10.5, bold=True, color_rgb=(20, 20, 20))
    r = p.add_run(text)
    set_run_font(r, "Calibri", 10.5, bold=False, color_rgb=(40, 40, 40))
    return p

def add_image_with_caption(filename, caption_text, width_inches=6.8):
    filepath = os.path.join(SCREENSHOTS_DIR, filename)
    if os.path.exists(filepath):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(filepath, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        set_run_font(r_cap, "Calibri", 9.5, italic=True, color_rgb=(80, 80, 80))
    else:
        p_err = doc.add_paragraph()
        r = p_err.add_run(f"[Screenshot {filename} missing]")
        set_run_font(r, "Calibri", 10, bold=True, color_rgb=(200, 0, 0))

# ----------------- BUILD DOCUMENT CONTENT -----------------

add_header_block()

add_callout_box(
    "A university Central Identity and User Management Platform developed as independent microservices "
    "communicating through an API Gateway with MongoDB Atlas persistence. Features two user roles: User (students/staff managing own profile) "
    "and Admin (administrators managing accounts). Strictly enforces mutual exclusion: Users can NEVER access Admin APIs, "
    "and Admins can NEVER access User APIs.",
    bold_prefix="Real-Life Scenario & System Architecture: "
)

# ARCHITECTURE SUMMARY TABLE
table = doc.add_table(rows=6, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Microservice", "Port", "Endpoints", "Role / Responsibilities"]
data = [
    ["API Gateway", "4000", "All routes (/) via reverse proxy", "Single entry point, JWT verification, role-based route guard"],
    ["Registration", "5001", "POST /register/userregister", "Validates input, checks duplicate email, bcrypt password hash, MongoDB store"],
    ["Login (Auth)", "5002", "POST /auth/login", "Validates email/password/role, issues signed 24h JWT token"],
    ["Admin Service", "5003", "GET /admin/searchuser, viewalluser\nDELETE /admin/deluser", "Admin operations: search by name/email, view all users, delete user"],
    ["User Service", "5004", "GET /user/viewprofile\nPUT /user/updateprofile", "User operations: view authenticated profile, update profile details"]
]
for col_idx, text in enumerate(headers):
    cell = table.cell(0, col_idx)
    shading = parse_xml(r'<w:shd {} w:fill="003366"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, "Calibri", 10, bold=True, color_rgb=(255, 255, 255))

for row_idx, row_data in enumerate(data, start=1):
    for col_idx, text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        if row_idx % 2 == 1:
            shading = parse_xml(r'<w:shd {} w:fill="F7F9FA"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        set_run_font(r, "Calibri", 9.5, bold=(col_idx==0), color_rgb=(30, 30, 30))

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# 1. TASK 1
add_heading_1("Task 1 – Five Microservices Running")
add_body_p(
    "All five independent microservices (API Gateway on port 4000, Registration Service on port 5001, "
    "Login Service on port 5002, Admin Service on port 5003, and User Service on port 5004) started and running simultaneously. "
    "TCP port listener verification confirms each process is listening on its designated port.",
    bold_lead="Microservices Orchestration:"
)
add_image_with_caption(
    "task1_2_microservices_running.png",
    "Evidence 1: PowerShell Orchestrator - All 5 Microservices running simultaneously on ports 4000, 5001, 5002, 5003, and 5004."
)

# 2. TASK 2
add_heading_1("Task 2 – MongoDB Database Connection")
add_body_p(
    "Mongoose database connection module (dbconnect.js) implemented with connection pooling, retry options, "
    "and deployment ping command. Successful execution confirms active connection to MongoDB Atlas database (clusterdb).",
    bold_lead="Database Connectivity:"
)
add_image_with_caption(
    "task2_dbconnect.png",
    "Evidence 2: VS Code editor showing dbconnect.js and terminal verification showing successful MongoDB Atlas connection."
)

# 3. TASK 3
add_heading_1("Task 3 – MongoDB User Model")
add_body_p(
    "Mongoose User schema (user_schema.js) created according to specification with _id, name, email (unique, lowercase), "
    "password (bcrypt hashed), role (enum: ['admin', 'user']), phone, and automatic timestamps (createdAt, updatedAt).",
    bold_lead="Schema Definition:"
)
add_image_with_caption(
    "task3_user_model_code.png",
    "Evidence 3: VS Code editor - MongoDB User Model Schema (user_schema.js) with required fields, validation, and timestamps."
)

# 4. TASK 4 & 7
add_heading_1("Task 4 – API Gateway Routing and Authentication")
add_body_p(
    "API Gateway (api-gateway.js) implemented as the single entry point on port 4000 using http-proxy reverse routing. "
    "Enforces strict JWT token validation and role-based access control (RBAC): ordinary users cannot access /admin/*, "
    "and administrators cannot access /user/*.",
    bold_lead="Gateway Routing & Security Architecture:"
)
add_image_with_caption(
    "task4_7_apigateway_code.png",
    "Evidence 4: VS Code editor - API Gateway source code (api-gateway.js) displaying JWT verification, role guard middleware, and reverse proxy routes."
)

# 5. TASK 5A
add_heading_1("Task 5 – Successful User Registration")
add_body_p(
    "Client registers a new student account via POST http://localhost:4000/register/userregister through the API Gateway. "
    "Registration service validates input, securely hashes the password with bcrypt, stores the document in MongoDB, and returns Status 201 Created.",
    bold_lead="Registration Workflow:"
)
add_image_with_caption(
    "task5_1_register_user.png",
    "Evidence 5: Postman - POST http://localhost:4000/register/userregister successfully registering student account (Status 201 Created)."
)

# Extra: Admin Registration
add_body_p(
    "Admin account registered via POST http://localhost:4000/register/userregister with role 'admin' for subsequent administrative operations.",
    bold_lead="Admin Account Setup:"
)
add_image_with_caption(
    "task5_2_register_admin.png",
    "Evidence 5b: Postman - POST http://localhost:4000/register/userregister successfully registering administrator account (Status 201 Created)."
)

# 6. TASK 5B
add_heading_1("Task 5 – Duplicate Email Validation")
add_body_p(
    "Attempting to register with an email address that already exists in MongoDB is immediately rejected with Status 400 Bad Request "
    "and error message 'Duplicate email not accepted. A user with this email address already exists.'",
    bold_lead="Duplicate Prevention:"
)
add_image_with_caption(
    "task5_3_duplicate_email.png",
    "Evidence 6: Postman - Duplicate email rejected with Status 400 Bad Request ('Duplicate email not accepted')."
)

# 7. TASK 5C
add_heading_1("Task 5 – MongoDB User Record with Hashed Password")
add_body_p(
    "MongoDB Atlas collection (clusterdb.users) inspection showing the stored student and administrator documents. "
    "Passwords are never stored in plain text and are securely hashed with bcrypt ($2a$10$...).",
    bold_lead="MongoDB Persistence Verification:"
)
add_image_with_caption(
    "task5_4_mongodb_users.png",
    "Evidence 7: MongoDB Compass / Atlas Data View - Registered user documents showing bcrypt password hashes."
)

# 8. TASK 6A
add_heading_1("Task 6 – Successful User Login and JWT Generation")
add_body_p(
    "User logs in via POST http://localhost:4000/auth/login through the API Gateway with valid email, password, and role. "
    "Login service verifies credentials against MongoDB and returns Status 200 OK along with a signed 24-hour JWT token.",
    bold_lead="Authentication Workflow:"
)
add_image_with_caption(
    "task6_1_user_login.png",
    "Evidence 8: Postman - POST http://localhost:4000/auth/login returning Status 200 OK and signed JWT token for student user."
)

# Extra: Admin Login
add_body_p(
    "Administrator logs in via POST http://localhost:4000/auth/login returning Status 200 OK and signed Admin JWT token.",
    bold_lead="Admin Authentication:"
)
add_image_with_caption(
    "task6_2_admin_login.png",
    "Evidence 8b: Postman - POST http://localhost:4000/auth/login returning Status 200 OK and signed Admin JWT token."
)

# 9. TASK 6B
add_heading_1("Task 6 – Invalid Password")
add_body_p(
    "Login attempt with an incorrect password is rejected with Status 401 Unauthorized ('Invalid credentials: Password does not match').",
    bold_lead="Password Verification:"
)
add_image_with_caption(
    "task6_3_invalid_password.png",
    "Evidence 9: Postman - Incorrect password rejected with Status 401 Unauthorized ('Password does not match')."
)

# 10. TASK 6C
add_heading_1("Task 6 – Invalid Role")
add_body_p(
    "Login attempt where the requested role does not match the account's registered role in MongoDB is rejected with Status 403 Forbidden ('Invalid role: Access denied. Account is registered as user, but login requested as admin').",
    bold_lead="Role Verification:"
)
add_image_with_caption(
    "task6_4_invalid_role.png",
    "Evidence 10: Postman - Student account attempting admin login rejected with Status 403 Forbidden ('Role mismatch')."
)

# 11. TASK 7 (Summary)
add_heading_1("Task 7 – JWT Authentication and Role-Based API Routing")
add_body_p(
    "JWT Authentication and RBAC guards implemented at the API Gateway level. Requests to /admin/* are validated to require decoded.role === 'admin'. "
    "Requests to /user/* are validated to require decoded.role === 'user'. Any violation generates immediate 401/403 responses before reaching the microservices.",
    bold_lead="Role-Based Security Enforcement:"
)

# 12. TASK 8A
add_heading_1("Task 8 – Admin Search User (User Found)")
add_body_p(
    "Administrator searches for a student user by query parameter (GET http://localhost:4000/admin/searchuser?email=...) with Admin Bearer Token. "
    "Admin service queries MongoDB and returns Status 200 OK with the matched user document.",
    bold_lead="Admin Search API:"
)
add_image_with_caption(
    "task8_2_admin_search_found.png",
    "Evidence 11: Postman - GET http://localhost:4000/admin/searchuser?email=... returning matching user document (Status 200 OK)."
)

# 13. TASK 8B
add_heading_1("Task 8 – Admin Search User (User Not Found)")
add_body_p(
    "Searching for a non-existent user returns Status 404 Not Found with error message 'User not found with the specified search criteria.'",
    bold_lead="Admin Search Handling:"
)
add_image_with_caption(
    "task8_3_admin_search_notfound.png",
    "Evidence 12: Postman - GET http://localhost:4000/admin/searchuser?email=ghost... returning Status 404 Not Found."
)

# 14. TASK 8 - VIEW ALL
add_heading_1("Task 8 – View All Users")
add_body_p(
    "Administrator retrieves all registered users via GET http://localhost:4000/admin/viewalluser with Admin Bearer Token. "
    "Admin service returns Status 200 OK with all user accounts (passwords excluded for security).",
    bold_lead="Admin View All API:"
)
add_image_with_caption(
    "task8_1_admin_viewalluser.png",
    "Evidence 13: Postman - GET http://localhost:4000/admin/viewalluser with Admin Token returning all registered users (Status 200 OK)."
)

# 15. TASK 8 - DELETE USER
add_heading_1("Task 8 – Delete User")
add_body_p(
    "Administrator deletes a user account via DELETE http://localhost:4000/admin/deluser?email=... with Admin Bearer Token. "
    "Admin service removes the user document from MongoDB and returns Status 200 OK.",
    bold_lead="Admin Delete API:"
)
add_image_with_caption(
    "task8_4_admin_delete_user.png",
    "Evidence 14: Postman - DELETE http://localhost:4000/admin/deluser?email=... successfully deleting user (Status 200 OK)."
)

# 16. TASK 9A
add_heading_1("Task 9 – View User Profile Before Update")
add_body_p(
    "Student views personal profile via GET http://localhost:4000/user/viewprofile using User Bearer Token. "
    "API Gateway passes identity to User Microservice, returning Status 200 OK with initial profile details.",
    bold_lead="User Profile View:"
)
add_image_with_caption(
    "task9_1_view_profile_before.png",
    "Evidence 15: Postman - GET http://localhost:4000/user/viewprofile retrieving original user profile (Status 200 OK)."
)

# 17. TASK 9B
add_heading_1("Task 9 – Update User Profile")
add_body_p(
    "Student updates name and phone number via PUT http://localhost:4000/user/updateprofile using User Bearer Token. "
    "User Microservice updates the document in MongoDB and returns Status 200 OK with updated profile information.",
    bold_lead="User Profile Update:"
)
add_image_with_caption(
    "task9_2_update_profile.png",
    "Evidence 16: Postman - PUT http://localhost:4000/user/updateprofile updating user name and phone (Status 200 OK)."
)

# Extra: View Profile After Update
add_body_p(
    "Student retrieves profile again via GET http://localhost:4000/user/viewprofile using User Bearer Token, confirming updated profile details (name and phone) are returned with Status 200 OK.",
    bold_lead="User Profile Verification After Update:"
)
add_image_with_caption(
    "task9_3_view_profile_after.png",
    "Evidence 16b: Postman - GET http://localhost:4000/user/viewprofile verifying updated profile data (Status 200 OK)."
)

# 18. TASK 9C
add_heading_1("Task 9 – MongoDB User Profile After Update")
add_body_p(
    "MongoDB Atlas collection inspection confirming the profile was modified in the database: name changed to 'Goat Ronaldo (Updated Profile)', "
    "phone changed to '+85599887766', and updatedAt timestamp automatically refreshed.",
    bold_lead="Database Verification After Update:"
)
add_image_with_caption(
    "task9_4_mongodb_profile_updated.png",
    "Evidence 17: MongoDB Compass / Atlas Data View - User profile in clusterdb.users showing updated name, phone, and refreshed timestamp."
)

# 19. TASK 10A
add_heading_1("Task 10(a) – Access Admin API Without Token")
add_body_p(
    "Client attempts to access protected route GET http://localhost:4000/admin/viewalluser without an Authorization header. "
    "API Gateway intercepts and immediately rejects the request with Status 401 Unauthorized ('Access Denied: No token provided').",
    bold_lead="Security Test 10(a):"
)
add_image_with_caption(
    "task10_a_without_token.png",
    "Evidence 18: Postman - Protected endpoint accessed without Authorization header rejected with Status 401 Unauthorized."
)

# 20. TASK 10B
add_heading_1("Task 10(b) – Access API Using Invalid JWT")
add_body_p(
    "Client attempts to access protected route using a fabricated or invalid token (Authorization: Bearer invalid.token.12345). "
    "API Gateway verifies token signature, detects malformed token, and rejects with Status 403 Forbidden ('Access Denied: Invalid token').",
    bold_lead="Security Test 10(b):"
)
add_image_with_caption(
    "task10_b_wrong_token.png",
    "Evidence 19: Postman - Protected endpoint accessed with invalid JWT rejected with Status 403 Forbidden."
)

# 21. TASK 10C
add_heading_1("Task 10(c) – Admin Token Attempting to Access User API")
add_body_p(
    "Client uses an Administrator JWT token to access the User-only endpoint GET http://localhost:4000/user/viewprofile. "
    "API Gateway enforces role mutual exclusion and rejects the request with Status 403 Forbidden ('Access Denied: User privileges required. Administrators are not permitted to access User APIs').",
    bold_lead="Security Test 10(c):"
)
add_image_with_caption(
    "task10_c_admin_access_user_forbidden.png",
    "Evidence 20: Postman - Administrator token attempting to access User API rejected with Status 403 Forbidden."
)

# 22. TASK 10D
add_heading_1("Task 10(d) – User Token Attempting to Access Admin API")
add_body_p(
    "Client uses an Ordinary User (Student) JWT token to access the Administrative endpoint GET http://localhost:4000/admin/viewalluser. "
    "API Gateway enforces role mutual exclusion and rejects the request with Status 403 Forbidden ('Access Denied: Admin privileges required. Ordinary users are not permitted to access Admin APIs').",
    bold_lead="Security Test 10(d):"
)
add_image_with_caption(
    "task10_d_user_access_admin_forbidden.png",
    "Evidence 21: Postman - Ordinary user token attempting to access Admin API rejected with Status 403 Forbidden."
)

# 23. TASK 11
add_heading_1("Task 11 – Public GitHub Repository")
add_body_p(
    "The complete source code including all 5 microservices, schemas, environment templates, and automated test scripts "
    "is published in a public GitHub repository. Sensitive files (.env, node_modules) are properly excluded via .gitignore.",
    bold_lead="Public Repository Link:"
)

p_link = doc.add_paragraph()
r_link_label = p_link.add_run("Public GitHub Repository URL: ")
set_run_font(r_link_label, "Calibri", 11, bold=True, color_rgb=(0, 51, 102))
r_url = p_link.add_run("https://github.com/goatronaldo/cld_assignment1_usermanagement.git")
set_run_font(r_url, "Calibri", 11, bold=True, color_rgb=(0, 102, 204))

add_image_with_caption(
    "task11_github_repo.png",
    "Evidence 22: GitHub Web View - Public repository goatronaldo/cld_assignment1_usermanagement showing the 5 microservices and project files."
)

# Save DOCX files
doc.save(DOCX_OUT_ROOT)
doc.save(DOCX_OUT_PROJ)
print(f"Saved DOCX successfully to: {DOCX_OUT_ROOT}")
print(f"Saved DOCX copy to: {DOCX_OUT_PROJ}")

# Convert to PDF using win32com
try:
    import win32com.client
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    docx_obj = word.Documents.Open(DOCX_OUT_ROOT)
    docx_obj.SaveAs(PDF_OUT_ROOT, FileFormat=17)
    docx_obj.SaveAs(PDF_OUT_PROJ, FileFormat=17)
    docx_obj.Close()
    word.Quit()
    print(f"Saved PDF successfully to: {PDF_OUT_ROOT}")
    print(f"Saved PDF copy to: {PDF_OUT_PROJ}")
except Exception as e:
    print(f"Word COM conversion note: {e}")
